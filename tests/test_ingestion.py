import io
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from lockstep.core.exceptions import IngestionError
from lockstep.services.ingestion import (
    clerical_key,
    parse_amount,
    parse_bool,
    parse_csv,
    parse_date,
    parse_excel,
    resolve_columns,
    to_canonical_rows,
)


def test_resolve_columns_recognizes_tally_style_headers():
    columns = ["Invoice No", "Invoice Date", "Supplier GSTIN", "Supplier Name", "Taxable Value"]
    resolved = resolve_columns(columns)

    assert resolved["invoice_number"] == "Invoice No"
    assert resolved["gstin"] == "Supplier GSTIN"
    assert resolved["taxable_value"] == "Taxable Value"


def test_resolve_columns_recognizes_gstr2b_style_headers():
    columns = [
        "GSTIN of Supplier", "Trade/Legal Name", "Invoice Number", "Invoice Date", "Taxable Value",
    ]
    resolved = resolve_columns(columns)

    assert resolved["gstin"] == "GSTIN of Supplier"
    assert resolved["vendor_name"] == "Trade/Legal Name"
    assert resolved["invoice_number"] == "Invoice Number"


def test_resolve_columns_recognizes_tally_igst_export_headers():
    """IGST and CGST/SGST arrive as separate exports with different column sets."""
    columns = ["Date", "Particulars", "Voucher No.", "Gross Total", "IGST Input"]
    resolved = resolve_columns(columns)

    assert resolved["invoice_number"] == "Voucher No."
    assert resolved["taxable_value"] == "Gross Total"
    assert resolved["integrated_tax"] == "IGST Input"


def test_resolve_columns_finds_the_filing_date():
    """The most valuable column in the 2B file — it must never be dropped."""
    resolved = resolve_columns(["GSTIN of Supplier", "GSTR-1/IFF/GSTR-5 Filing Date"])
    assert resolved["filing_date"] == "GSTR-1/IFF/GSTR-5 Filing Date"


def test_parse_csv_raises_when_required_column_missing():
    with pytest.raises(IngestionError):
        parse_csv(b"Foo,Bar\n1,2\n")


def test_parse_csv_round_trips_rows():
    content = b"Invoice No,Supplier GSTIN,Taxable Value\nINV-1,29ABCDE1234F1Z5,1000\n"
    rows, columns = parse_csv(content)

    assert rows == [
        {"Invoice No": "INV-1", "Supplier GSTIN": "29ABCDE1234F1Z5", "Taxable Value": "1000"}
    ]
    assert columns["invoice_number"] == "Invoice No"


def test_parse_csv_skips_a_multi_row_header():
    """The portal's GSTR-2B download buries the real header a few rows down."""
    content = (
        b"Goods and Services Tax - GSTR-2B,,,\n"
        b"Taxpayer GSTIN,29ZZZZZ9999Z1Z9,,\n"
        b"GSTIN of Supplier,Invoice Number,Taxable Value,Integrated Tax\n"
        b"29ABCDE1234F1Z5,INV-1,10000,1800\n"
    )
    rows, columns = parse_csv(content)

    assert len(rows) == 1
    assert columns["gstin"] == "GSTIN of Supplier"
    assert rows[0]["Invoice Number"] == "INV-1"


def test_parse_csv_drops_trailing_total_rows():
    content = (
        b"Invoice No,Supplier GSTIN,Taxable Value\n"
        b"INV-1,29ABCDE1234F1Z5,1000\n"
        b",,5000\n"
    )
    rows, _ = parse_csv(content)
    assert len(rows) == 1


def test_parse_amount_strips_currency_and_commas():
    assert parse_amount("₹1,20,400.00") == Decimal("120400.00")
    assert parse_amount("") == Decimal("0")


def test_parse_amount_returns_decimal_not_float():
    """Money must never round-trip through a float — it disagrees with NUMERIC(14,2)."""
    assert isinstance(parse_amount("0.1"), Decimal)
    assert parse_amount("0.1") + parse_amount("0.2") == Decimal("0.30")


def test_parse_date_handles_dd_mm_yyyy():
    assert parse_date("05-05-2025") == date(2025, 5, 5)


def test_parse_bool_reads_portal_yes_no():
    assert parse_bool("Y") is True
    assert parse_bool("No") is False
    assert parse_bool("") is None


def test_clerical_key_folds_ocr_confusions_and_leading_zeros():
    assert clerical_key("INV-007") == clerical_key("INV-7")
    assert clerical_key("INV-O12") == clerical_key("INV-012")
    assert clerical_key("inv 1") == clerical_key("INV-1")


def test_to_canonical_rows_splits_tax_heads():
    content = (
        b"GSTIN of Supplier,Invoice Number,Taxable Value,Central Tax,State/UT Tax\n"
        b"29ABCDE1234F1Z5,INV-1,1000,90,90\n"
    )
    rows, columns = parse_csv(content)
    row = to_canonical_rows(rows, columns)[0]

    assert row.cgst == Decimal("90.00")
    assert row.sgst == Decimal("90.00")
    assert row.total_tax == Decimal("180.00")
    assert row.taxable_value == Decimal("1000.00")


def test_to_canonical_rows_keeps_filing_date_and_itc_eligibility():
    content = (
        b"GSTIN of Supplier,Invoice Number,Taxable Value,GSTR-1/IFF/GSTR-5 Filing Date,"
        b"ITC Availability,Reason,Supply Attract Reverse Charge\n"
        b"29ABCDE1234F1Z5,INV-1,1000,17-09-2026,No,POS and supplier state are same,Y\n"
    )
    rows, columns = parse_csv(content)
    row = to_canonical_rows(rows, columns)[0]

    assert row.supplier_filed_at == date(2026, 9, 17)
    assert row.itc_available is False
    assert row.itc_reason == "POS and supplier state are same"
    assert row.is_reverse_charge is True


def test_itc_amount_falls_back_to_taxable_value_without_tax_columns():
    content = b"Invoice No,Supplier GSTIN,Taxable Value\nINV-1,29ABCDE1234F1Z5,1000\n"
    rows, columns = parse_csv(content)
    assert to_canonical_rows(rows, columns)[0].itc_amount == Decimal("1000.00")


def test_taxable_value_is_not_confused_with_invoice_value():
    """A real GSTR-2B export carries both columns — 'Invoice Value(₹)' includes tax,
    'Taxable Value (₹)' doesn't, and only the latter is the canonical field."""
    columns = [
        "GSTIN of supplier", "Invoice number", "Invoice Value(₹)", "Taxable Value (₹)",
    ]
    resolved = resolve_columns(columns)
    assert resolved["taxable_value"] == "Taxable Value (₹)"


def test_parse_amount_ignores_currency_suffix_in_header_matching():
    content = (
        b"Invoice number,GSTIN of supplier,Taxable Value (\xe2\x82\xb9)\n"
        b"INV-1,29ABCDE1234F1Z5,1000\n"
    )
    rows, columns = parse_csv(content)
    assert columns["taxable_value"] == "Taxable Value (₹)"
    assert to_canonical_rows(rows, columns)[0].taxable_value == Decimal("1000.00")


def _two_row_header_workbook() -> bytes:
    """Build the real GSTR-2B portal shape in memory: a parent header row with
    spanning labels ('Invoice Details', 'Tax Amount') and a child sub-header row
    directly below it (mirrors the vendor-supplied 'GST - Igst.xlsx' fixture)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GSTR 2B"
    ws.append(["Goods and Services Tax - GSTR-2B"])
    ws.append([])
    ws.append(["Taxable inward supplies received from registered persons"])
    ws.append([
        "GSTIN of supplier", "Trade/Legal name", "Invoice Details", None, None, None,
        "Place of supply", "Supply Attract Reverse Charge", "Taxable Value (₹)",
        "Tax Amount", None, None, None,
        "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date",
        "ITC Availability", "Reason",
    ])
    ws.append([
        None, None, "Invoice number", "Invoice type", "Invoice Date", "Invoice Value(₹)",
        None, None, None, "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)",
        None, None, None, None,
    ])
    ws.append([
        "29ABCDE1234F1Z5", "Sharma Traders", "INV-1", "Regular", "2026-08-05", 118000,
        "29-Karnataka", "N", 100000, 18000, 0, 0, 0, "082026", "2026-09-10", "Yes", "",
    ])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parse_excel_merges_a_two_row_portal_header():
    rows, columns = parse_excel(_two_row_header_workbook())

    assert len(rows) == 1
    assert columns["invoice_number"] == "Invoice number"
    assert columns["gstin"] == "GSTIN of supplier"
    assert columns["taxable_value"] == "Taxable Value (₹)"
    assert columns["integrated_tax"] == "Integrated Tax(₹)"
    row = to_canonical_rows(rows, columns)[0]
    assert row.invoice_number == "INV-1"
    assert row.taxable_value == Decimal("100000.00")
    assert row.igst == Decimal("18000.00")


def test_parse_excel_picks_the_named_sheet_out_of_a_multi_sheet_workbook():
    """A single Tally export can bundle an IGST ledger, a CGST/SGST ledger and the
    GSTR-2B download as three sheets of one workbook."""
    wb = openpyxl.load_workbook(io.BytesIO(_two_row_header_workbook()))
    wb.create_sheet("Tally -IGST")
    buffer = io.BytesIO()
    wb.save(buffer)

    rows, columns = parse_excel(buffer.getvalue(), hint="gstr2b")
    assert len(rows) == 1
    assert columns["invoice_number"] == "Invoice number"


def test_normalize_gstin_rejects_the_wrong_length():
    """An overlong GSTIN is a data-entry error or the wrong column mapped — never
    truncate it into a different, wrong-but-valid-looking one."""
    from lockstep.services.ingestion import normalize_gstin

    assert normalize_gstin("29ABCDE1234F1Z5") == "29ABCDE1234F1Z5"
    assert normalize_gstin("29ABCDE1234F1Z5EXTRA") == ""
    assert normalize_gstin("29ABC") == ""
    assert normalize_gstin("") == ""


def test_vendor_email_and_phone_are_captured_when_present():
    content = (
        b"Invoice No,Supplier GSTIN,Taxable Value,Vendor Email,Contact No\n"
        b"INV-1,29ABCDE1234F1Z5,1000,accounts@example.com,9876543210\n"
    )
    rows, columns = parse_csv(content)
    row = to_canonical_rows(rows, columns)[0]
    assert row.vendor_email == "accounts@example.com"
    assert row.vendor_phone == "9876543210"
